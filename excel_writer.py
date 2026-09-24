from datetime import datetime

from openpyxl import Workbook, load_workbook

COLUMNS = [
    "FROM_DATE", "TO_DATE", "STATEMENT_LANGUAGE", "FULL_NAME", "FINANSIAL_INSTITUTION",
    "AMOUNT", "DETAILS", "OPERATION_DATE", "TRANSACTION_TYPE", "INSERT_DATE", "CARD_NUMBER",
    "ST_CREATION_DATE", "ST_MODIFIED_DATE", "ST_SUBJECT", "ST_AUTHOR", "ST_TITLE", "ST_PRODUCER",
]


def write_to_excel(st, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "statements"
        ws.append(COLUMNS)

    now = datetime.now().isoformat()
    for t in st.transactions:
        ws.append([
            st.from_date, st.to_date, st.language, st.full_name, st.bank,
            t.amount, t.details, t.operation_date, t.transaction_type, now, st.card_number,
            st.meta["creation_date"], st.meta["modification_date"], st.meta["subject"],
            st.meta["author"], st.meta["title"], st.meta["producer"],
        ])

    wb.save(path)
